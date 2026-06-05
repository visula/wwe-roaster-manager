import openpyxl
import json
import csv

excel_file = r'c:\Users\svisu\Downloads\Online Version WWE 2k26 Universe Mode Gameplay Plan.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)

wrestlers = []
matches_by_show = {}

# Extract wrestlers from WWE 2K26 Universal sheet
try:
    ws = wb['WWE 2K26 Universal']
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and str(row[0]).lower() != 'name':
            try:
                name = str(row[0]).strip() if row[0] else None
                overall = str(row[2]).strip() if row[2] else 'Unassigned'
                alignment = str(row[3]).strip() if row[3] else 'Active'
                show = str(row[4]).strip() if row[4] else 'RAW'
                
                if name and name.lower() not in ['name', 'none', '']:
                    wrestlers.append({
                        'name': name,
                        'show': show,
                        'division': overall,
                        'status': alignment if alignment else 'Active'
                    })
            except:
                pass
except Exception as e:
    print(f"Error reading wrestlers: {e}")

# Extract matches from various match sheets
match_sheets = ['RAW Matches', 'Smackdown Matches', 'NXT Matches', 'TNA Matches', 'AAA Matches', 'AEW Matches']

for sheet_name in match_sheets:
    try:
        ws = wb[sheet_name]
        show = sheet_name.replace(' Matches', '')
        matches_by_show[show] = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and str(row[0]).lower() not in ['date', 'none', '']:
                try:
                    match_data = {
                        'show': show,
                        'date': str(row[0]) if row[0] else '',
                        'type': str(row[1]).strip() if row[1] else 'Singles',
                        'participant1': str(row[2]).strip() if row[2] else '',
                        'participant2': str(row[3]).strip() if row[3] else '',
                        'participant3': str(row[4]).strip() if row[4] else '',
                        'participant4': str(row[5]).strip() if row[5] else '',
                        'result': str(row[6]).strip() if row[6] else 'Pending',
                        'winner': str(row[7]).strip() if row[7] else ''
                    }
                    if match_data['participant1']:
                        matches_by_show[show].append(match_data)
                except:
                    pass
    except:
        pass

# Extract championships
championships = []
try:
    ws = wb['Championships']
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and str(row[0]).lower() not in ['name', 'none', '']:
            try:
                champ_data = {
                    'name': str(row[2]).strip() if row[2] else '',
                    'show': str(row[0]).strip() if row[0] else 'RAW',
                    'holder': str(row[3]).strip() if row[3] else 'Vacant',
                    'debutDate': ''
                }
                if champ_data['name']:
                    championships.append(champ_data)
            except:
                pass
except:
    pass

# Save to CSV files
output_wrestlers = r'f:\Work\Fun-Projects\data\wrestlers-import.csv'
with open(output_wrestlers, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=['name', 'show', 'division', 'status'])
    writer.writeheader()
    writer.writerows(wrestlers)

output_champs = r'f:\Work\Fun-Projects\data\championships-import.csv'
with open(output_champs, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=['name', 'show', 'holder', 'debutDate'])
    writer.writeheader()
    writer.writerows(championships)

print(f"✅ Extracted {len(wrestlers)} wrestlers")
print(f"✅ Extracted {len(championships)} championships")
print(f"✅ Extracted {sum(len(m) for m in matches_by_show.values())} matches")
print("\nData exported to:")
print(f"  - {output_wrestlers}")
print(f"  - {output_champs}")

# Print sample data
print("\n=== Sample Wrestlers ===")
for w in wrestlers[:5]:
    print(f"  {w['name']} ({w['show']})")

print("\n=== Sample Championships ===")
for c in championships[:5]:
    print(f"  {c['name']} ({c['show']}) - {c['holder']}")
