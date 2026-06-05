import openpyxl
import csv
from pathlib import Path

excel_file = r'c:\Users\svisu\Downloads\Online Version WWE 2k26 Universe Mode Gameplay Plan.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)
ws = wb['WWE Universal']

# Column mapping
SHOW_COLUMNS = {
    'N': 'RAW',
    'O': 'AAA',
    'P': 'TNA',
    'Q': 'NXT',
    'R': 'AEW',
    'S': 'SmackDown'
}

# Extract data
wrestlers = []
seen = set()

for row_num in range(2, ws.max_row + 1):
    name = ws.cell(row=row_num, column=1).value
    if not name or name.lower() == 'name':
        continue
    
    name = str(name).strip()
    
    # Skip duplicates in this extraction
    if name in seen:
        continue
    seen.add(name)
    
    gender = str(ws.cell(row=row_num, column=3).value or '').strip() or 'Unknown'
    overall = ws.cell(row=row_num, column=4).value
    alignment = str(ws.cell(row=row_num, column=5).value or '').strip() or 'Active'
    primary_show = str(ws.cell(row=row_num, column=6).value or '').strip() or 'Unassigned'
    tag_team = str(ws.cell(row=row_num, column=7).value or '').strip() or ''
    tag_partners = str(ws.cell(row=row_num, column=12).value or '').strip() or ''
    
    # Get all shows where wrestler appears ('Y' in columns N-S)
    shows = []
    for col_letter, show_name in SHOW_COLUMNS.items():
        cell_value = str(ws.cell(row=row_num, column=openpyxl.utils.column_index_from_string(col_letter)).value or '').strip().upper()
        if cell_value == 'Y':
            shows.append(show_name)
    
    # If no shows mapped but primary show exists, use that
    if not shows and primary_show and primary_show != 'Unassigned':
        shows = [primary_show]
    
    # If still no shows, default to primary show or 'Unassigned'
    if not shows:
        shows = [primary_show] if primary_show else ['Unassigned']
    
    # Create one record per show (since our current DB structure requires a single show)
    for show in shows:
        wrestlers.append({
            'name': name,
            'gender': gender,
            'overall': overall or '',
            'alignment': alignment,
            'show': show,
            'division': 'Unassigned',
            'status': 'Active',
            'tag_team': tag_team,
            'tag_partners': tag_partners
        })

print(f"Extracted {len(wrestlers)} wrestler-show assignments")

# Show breakdown
show_counts = {}
for w in wrestlers:
    show = w['show']
    show_counts[show] = show_counts.get(show, 0) + 1

print("\nBreakdown by show:")
for show, count in sorted(show_counts.items(), key=lambda x: -x[1]):
    print(f"  {show}: {count}")

# Write to CSV
output_file = Path('f:\\Work\\Fun-Projects\\data\\wrestlers-universal.csv')
output_file.parent.mkdir(parents=True, exist_ok=True)

with open(output_file, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=['name', 'gender', 'overall', 'alignment', 'show', 'division', 'status', 'tag_team', 'tag_partners'])
    writer.writeheader()
    writer.writerows(wrestlers)

print(f"\n✅ Extracted to {output_file}")
